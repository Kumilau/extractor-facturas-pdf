"""Exporta las facturas extraídas a un Excel con formato."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .extraccion import Factura
from .rut import con_puntos

AZUL = "1F4E79"
AMARILLO = "FFE699"
VERDE = "C6EFCE"
ROJO = "F8CBAD"
GRIS = "D9D9D9"
PESOS = '"$"#,##0'
FECHA = 'dd-mm-yyyy'

_lado = Side(style="thin", color="BFBFBF")
BORDE = Border(left=_lado, right=_lado, top=_lado, bottom=_lado)


def _relleno(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


def _titulo(ws: Worksheet, titulo: str, subtitulo: str) -> None:
    ws["A1"] = titulo
    ws["A1"].font = Font(size=16, bold=True, color=AZUL)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(italic=True, color="595959")


def _tabla(ws: Worksheet, encabezados: list[str], filas: list[list], formatos: dict[int, str],
           fila_inicio: int = 4) -> int:
    for c, nombre in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_inicio, column=c, value=nombre)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = _relleno(AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = BORDE
    for r, fila in enumerate(filas, start=fila_inicio + 1):
        for c, valor in enumerate(fila, start=1):
            celda = ws.cell(row=r, column=c, value=valor)
            celda.border = BORDE
            if c in formatos:
                celda.number_format = formatos[c]
    ws.freeze_panes = ws.cell(row=fila_inicio + 1, column=1)
    if filas:
        ws.auto_filter.ref = f"A{fila_inicio}:{get_column_letter(len(encabezados))}{fila_inicio + len(filas)}"
    return fila_inicio + len(filas)


def _fila_total(ws: Worksheet, fila: int, primera: int, columnas_suma: list[int], formato: str) -> None:
    ws.cell(row=fila, column=1, value="TOTAL")
    for c in range(1, ws.max_column + 1):
        celda = ws.cell(row=fila, column=c)
        celda.fill = _relleno(AMARILLO)
        celda.font = Font(bold=True)
        celda.border = BORDE
    for c in columnas_suma:
        letra = get_column_letter(c)
        celda = ws.cell(row=fila, column=c, value=f"=SUM({letra}{primera}:{letra}{fila - 1})")
        celda.number_format = formato


def _anchos(ws: Worksheet, anchos: list[int]) -> None:
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def generar_excel(facturas: list[Factura], salida: Path) -> Path:
    salida = Path(salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ok = [f for f in facturas if f.estado == "OK"]
    revisar = [f for f in facturas if f.estado == "Revisar"]
    no_rec = [f for f in facturas if f.estado == "No reconocido"]
    leidas = [f for f in facturas if f.estado != "No reconocido"]

    # Hoja 1: una fila por factura
    ws = wb.active
    ws.title = "Facturas"
    _titulo(ws, "Facturas extraídas",
            f"{len(facturas)} PDF procesados · {len(ok)} OK · {len(revisar)} para revisar · "
            f"{len(no_rec)} no reconocidos")
    encabezados = ["Archivo", "Proveedor", "RUT proveedor", "Folio", "Fecha", "Neto", "IVA", "Total",
                   "N° ítems", "Estado"]
    filas = [[f.archivo, f.proveedor, con_puntos(f.rut_proveedor), f.folio, f.fecha, f.neto, f.iva, f.total,
              len(f.items), f.estado] for f in leidas]
    ultima = _tabla(ws, encabezados, filas, {5: FECHA, 6: PESOS, 7: PESOS, 8: PESOS})
    for r in range(5, ultima + 1):
        estado = ws.cell(row=r, column=10)
        estado.fill = _relleno(VERDE if estado.value == "OK" else ROJO)
        estado.alignment = Alignment(horizontal="center")
    if filas:
        _fila_total(ws, ultima + 1, 5, [6, 7, 8], PESOS)
    _anchos(ws, [22, 34, 16, 10, 12, 14, 13, 14, 9, 12])

    # Hoja 2: detalle de productos
    ws = wb.create_sheet("Detalle")
    _titulo(ws, "Detalle de productos y servicios", "Una fila por línea de cada factura")
    filas = [[f.archivo, f.proveedor, f.folio, i["descripcion"], i["cantidad"], i["precio_unitario"],
              i["total_linea"]] for f in leidas for i in f.items]
    ultima = _tabla(ws, ["Archivo", "Proveedor", "Folio", "Descripción", "Cantidad", "Precio unitario",
                         "Total línea"], filas, {5: '#,##0', 6: PESOS, 7: PESOS})
    if filas:
        _fila_total(ws, ultima + 1, 5, [7], PESOS)
    _anchos(ws, [22, 34, 10, 38, 11, 16, 15])

    # Hoja 3: totales por proveedor
    ws = wb.create_sheet("Por proveedor")
    _titulo(ws, "Compras por proveedor", "Solo facturas en estado OK")
    resumen: dict[tuple, list] = {}
    for f in ok:
        clave = (f.proveedor, f.rut_proveedor)
        actual = resumen.setdefault(clave, [0, 0, 0, 0])
        actual[0] += 1
        actual[1] += f.neto
        actual[2] += f.iva
        actual[3] += f.total
    filas = [[p, con_puntos(r), *v] for (p, r), v in sorted(resumen.items(), key=lambda kv: -kv[1][3])]
    ultima = _tabla(ws, ["Proveedor", "RUT", "Facturas", "Neto", "IVA", "Total"], filas,
                    {3: '#,##0', 4: PESOS, 5: PESOS, 6: PESOS})
    if filas:
        _fila_total(ws, ultima + 1, 5, [3, 4, 5, 6], PESOS)
        ws.cell(row=ultima + 1, column=3).number_format = '#,##0'
    _anchos(ws, [34, 16, 10, 15, 14, 15])

    # Hoja 4: observaciones
    ws = wb.create_sheet("Revisión")
    pendientes = revisar + no_rec
    _titulo(ws, "Documentos para revisar",
            "Todo en orden: no hay observaciones." if not pendientes
            else f"{len(pendientes)} documentos requieren revisión manual antes de contabilizarse.")
    filas = [[f.archivo, f.proveedor, f.folio, f.estado, obs] for f in pendientes for obs in f.observaciones]
    _tabla(ws, ["Archivo", "Proveedor", "Folio", "Estado", "Observación"], filas, {})
    if pendientes:
        ws.sheet_properties.tabColor = "C00000"
    _anchos(ws, [26, 34, 10, 15, 70])

    for hoja in wb.worksheets:
        hoja.page_setup.orientation = "landscape"
        hoja.page_setup.fitToWidth = 1
        hoja.page_setup.fitToHeight = 0
        hoja.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(salida)
    return salida

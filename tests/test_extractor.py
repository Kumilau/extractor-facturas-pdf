import sys
from pathlib import Path

import pytest
from datetime import date
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

from extractor_facturas import rut  # noqa: E402
from extractor_facturas.excel import generar_excel  # noqa: E402
from extractor_facturas.extraccion import extraer_carpeta, extraer_de_texto  # noqa: E402
from generar_facturas_ejemplo import factura_formato_a, factura_formato_b  # noqa: E402

TEXTO_A = """Embalajes del Sur SpA R.U.T.: 76.482.113-0
Giro: Servicios logísticos FACTURA ELECTRÓNICA
Av. Ejemplo 1234, Santiago N° 42875
Fecha de emisión: 19/09/2026
Señor(es): Comercial Logística Ejemplo Ltda.
RUT: 76.555.222-2
Cantidad Descripción Precio Unit. Total
39 Film stretch 50 cm 12.900 503.100
Monto Neto $ 503.100
IVA (19%) $ 95.589
Total $ 598.689"""
TABLAS_A = [[["Cantidad", "Descripción", "Precio Unit.", "Total"], ["39", "Film stretch 50 cm", "12.900", "503.100"]]]

TEXTO_B = """Factura Electrónica
Emisor: Combustibles Ruta 5 Ltda. - RUT 76.118.562-4
Folio: 2215
Fecha: 2026-01-21
Cliente: Comercial Logística Ejemplo Ltda. - RUT 76.555.222-2
Neto: 1.105.000
IVA 19%: 209.950
TOTAL A PAGAR: 1.314.950"""


@pytest.mark.parametrize("valor, esperado", [
    ("76.482.113-0", True), ("76482113-0", True), ("96.874.320-2", False),
    ("11.111.111-1", True), ("12.345.678-5", True), ("12.345.678-K", False),
])
def test_validacion_rut(valor, esperado):
    assert rut.es_valido(valor) is esperado


def test_formato_clasico():
    f = extraer_de_texto(TEXTO_A, TABLAS_A, "a.pdf")
    assert f.proveedor == "Embalajes del Sur SpA"
    assert f.rut_proveedor == "76482113-0"
    assert f.rut_cliente == "76555222-2"
    assert f.folio == 42875
    assert str(f.fecha) == "2026-09-19"
    assert (f.neto, f.iva, f.total) == (503100, 95589, 598689)
    assert f.items == [{"descripcion": "Film stretch 50 cm", "cantidad": 39,
                        "precio_unitario": 12900, "total_linea": 503100}]
    assert f.estado == "OK"


def test_formato_moderno():
    f = extraer_de_texto(TEXTO_B, [], "b.pdf")
    assert f.proveedor == "Combustibles Ruta 5 Ltda."
    assert f.folio == 2215
    assert (f.neto, f.iva, f.total) == (1105000, 209950, 1314950)
    assert f.estado == "OK"


def test_total_que_no_cuadra():
    f = extraer_de_texto(TEXTO_B.replace("1.314.950", "1.324.950"), [], "b.pdf")
    assert f.estado == "Revisar"
    assert any("+10.000" in o for o in f.observaciones)


def test_iva_incorrecto():
    f = extraer_de_texto(TEXTO_B.replace("209.950", "200.000").replace("1.314.950", "1.305.000"), [], "b.pdf")
    assert any("19%" in o for o in f.observaciones)


def test_documento_que_no_es_factura():
    f = extraer_de_texto("Cotización de servicios\nNo constituye factura.", [], "c.pdf")
    assert f.estado == "No reconocido"


def test_campo_faltante():
    f = extraer_de_texto(TEXTO_B.replace("Folio: 2215\n", ""), [], "b.pdf")
    assert "No se encontró el folio" in f.observaciones


def test_pdf_reales_y_excel(tmp_path):
    items = [(3, "Pallet de madera 120x100", 8500), (2, "Film stretch 50 cm", 12900)]
    factura_formato_a(tmp_path / "a.pdf", "Embalajes del Sur SpA", rut.formatear(76482113), 111,
                      date(2026, 5, 4), items)
    factura_formato_b(tmp_path / "b.pdf", "TecnoOficina Chile SpA", rut.formatear(77930441), 222,
                      date(2026, 5, 5), items)
    factura_formato_b(tmp_path / "c.pdf", "TecnoOficina Chile SpA", rut.formatear(77930441), 222,
                      date(2026, 5, 5), items)  # duplicada
    facturas = extraer_carpeta(tmp_path)
    neto = 3 * 8500 + 2 * 12900
    assert [f.neto for f in facturas] == [neto] * 3
    assert [len(f.items) for f in facturas] == [2, 2, 2]
    assert [f.estado for f in facturas] == ["OK", "OK", "Revisar"]
    assert "duplicada" in facturas[2].observaciones[0]

    salida = generar_excel(facturas, tmp_path / "salida.xlsx")
    wb = load_workbook(salida)
    assert wb.sheetnames == ["Facturas", "Detalle", "Por proveedor", "Revisión"]
    assert wb["Facturas"]["H5"].value == neto + round(neto * 0.19)
